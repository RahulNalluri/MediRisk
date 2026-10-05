from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "MediRisk_Demo_Handoff_Report.docx"


def set_run(run, size=None, bold=None, color=None):
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Arial"
        if level == 1:
            run.font.color.rgb = RGBColor(31, 78, 121)
        elif level == 2:
            run.font.color.rgb = RGBColor(54, 96, 146)
    return p


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run(r, bold=True)
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                set_run(r, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    doc.add_paragraph()
    return table


def add_image_if_exists(doc, rel_path, caption, width=6.1):
    path = ROOT / rel_path
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(9)


def dataset_rows():
    rows = []
    for rel, purpose in [
        ("data/diabetes.csv", "Initial Pima diabetes dataset"),
        ("data/diabetes_clean.csv", "Cleaned binary diabetes dataset"),
        ("data/diabetes_dataset.csv", "Large diabetes-type dataset"),
        ("data/diabetes_dataset_clean.csv", "Cleaned diabetes-type dataset"),
        ("data/heart_disease_uci.csv", "Original UCI heart disease dataset"),
        ("data/heart_clean.csv", "Cleaned heart disease model dataset"),
    ]:
        path = ROOT / rel
        if path.exists():
            df = pd.read_csv(path)
            rows.append([rel, f"{df.shape[0]} rows x {df.shape[1]} columns", purpose])
    return rows


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MediRisk Project Demo Handoff Report")
    set_run(r, size=22, bold=True, color=(31, 78, 121))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Medical Data Analysis, Disease Risk Prediction, RAG Chatbot, and Hospital Workflow")
    set_run(r, size=12, bold=True, color=(80, 80, 80))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Prepared for demo | Generated on {date.today().isoformat()}")

    add_heading(doc, "1. Executive Summary")
    add_para(
        doc,
        "MediRisk started as a medical data analysis and disease-risk visualization project and evolved into a complete AI-assisted healthcare demo system. The final system combines machine learning prediction, exploratory visual analytics, a RAG-based medical chatbot, explainable result cards, multilingual support, emergency safety handling, hospital login, role-based dashboards, patient registration, report storage, doctor summaries, and patient portal access.",
    )
    add_bullets(
        doc,
        [
            "Core prediction tasks: diabetes risk, diabetes type classification, and heart disease risk.",
            "Core AI assistant tasks: medical education, Indian diet guidance, local fallback answers, and session/patient-aware context.",
            "Core hospital workflow: super admin, hospital admin, doctor, receptionist, and patient roles.",
            "Safety stance: screening and education only; not a diagnosis or replacement for a qualified doctor.",
        ]
    )

    add_heading(doc, "2. Project Evolution From Scratch")
    add_table(
        doc,
        ["Stage", "Work Completed", "Output"],
        [
            ["Initial analysis", "Loaded diabetes and heart disease datasets, checked columns, missing values, distributions, and target balance.", "Clean datasets and baseline plots."],
            ["Visualization", "Created disease counts, glucose/BMI/age distributions, correlation heatmaps, risk zones, and factor-specific graphs.", "Plots stored in the plots folder."],
            ["Model training", "Trained Random Forest, Neural Network, and XGBoost variants; used SMOTE for diabetes imbalance handling.", "Saved .pkl model artifacts."],
            ["Diabetes type expansion", "Used a 70,000-row diabetes-type dataset and encoded categorical variables for four-class classification.", "diabetes2_xgb_model.pkl and label encoder."],
            ["Backend integration", "Created Flask APIs for prediction, chat, simple heart screening, reports, patients, analytics, and auth.", "backend/app.py and supporting modules."],
            ["RAG assistant", "Built ChromaDB index over medical knowledge files using SentenceTransformer embeddings and OpenRouter LLM responses.", "84 indexed knowledge chunks."],
            ["Hospital workflow", "Added SQLite schema, role-based permissions, dashboards, patient portal accounts, doctor summaries, audit logs, and analytics.", "Demo-ready hospital system."],
        ],
    )

    add_heading(doc, "3. Datasets Used")
    add_table(doc, ["Dataset", "Shape", "Purpose"], dataset_rows())
    add_para(
        doc,
        "The project used both small clinical-style datasets and a larger diabetes-type dataset. The 768-row diabetes dataset was useful for learning binary diabetes prediction and feature analysis. The 70,000-row dataset was used for diabetes type classification across Prediabetic, Type 1, Type 2, and Type 3/Type 3c-style groups. The UCI heart dataset was cleaned and one-hot encoded for heart disease prediction.",
    )

    add_heading(doc, "4. Jupyter Notebooks")
    add_table(
        doc,
        ["Notebook", "Main Role"],
        [
            ["MediRisk_Diabetes_Analysis.ipynb", "Binary diabetes EDA, baseline training, Random Forest, Neural Network, XGBoost, SMOTE, plots, and saved diabetes models."],
            ["Diabetes2_Analysis.ipynb", "Large diabetes-type analysis and four-class XGBoost classifier with 95.07% accuracy."],
            ["MediRisk_HeartDisease_Analysis..ipynb", "UCI heart disease cleaning, one-hot encoding, EDA, XGBoost training, and 92.93% accuracy result."],
            ["Clustering.ipynb", "KMeans clustering and risk-group visualization for diabetes and heart datasets."],
            ["KnowledgeBase.ipynb", "Generated medical knowledge files from dataset insights and guideline-style content for RAG."],
        ],
    )

    add_heading(doc, "5. Data Visualization Work")
    add_para(doc, "The plots folder contains visual evidence used during analysis and demo explanation. Important graph categories include:")
    add_bullets(
        doc,
        [
            "Diabetes: disease count, glucose distribution, BMI distribution, age distribution, insulin, blood pressure, diabetes pedigree function, feature importance, model comparison, confusion matrix, and risk zones.",
            "Diabetes type: type distribution, glucose by type, age by type, BMI by type, correlation, causes, genetic/lifestyle/smoking/alcohol factors, and cluster groups.",
            "Heart disease: disease count, age, gender, chest pain, cholesterol, blood pressure, exercise pain, blocked vessels, correlation heatmap, KDE, and risk zones.",
        ],
    )
    add_image_if_exists(doc, "plots/diabetes_model_comparison.png", "Diabetes model comparison used during model selection.")
    add_image_if_exists(doc, "plots/d2_01_type_distribution.png", "Diabetes-type distribution from the 70,000-row dataset.")
    add_image_if_exists(doc, "plots/hd_09_correlation.png", "Heart disease correlation analysis.")
    add_image_if_exists(doc, "plots/hd_08_blocked_vessels.png", "Blocked vessels as a strong heart disease risk factor.")

    add_heading(doc, "6. Model Training Strategy")
    add_para(doc, "The modeling strategy was iterative: start with data cleaning and EDA, train baseline models, compare results, improve class balance where needed, and save the strongest practical models for backend use.")
    add_table(
        doc,
        ["Task", "Models Tried / Used", "Final or Notable Result"],
        [
            ["Binary diabetes prediction", "Random Forest, Neural Network, XGBoost, XGBoost + SMOTE, tuned XGBoost", "Random Forest about 75.97%; Neural Network about 70.13%; XGBoost about 72.08%; XGBoost + SMOTE about 72.73%. Saved final XGBoost and RF artifacts."],
            ["Diabetes type classification", "XGBoost multi-class classifier with label encoder", "95.07% accuracy on 14,000 test samples. Recall: Prediabetic 100%, Type 1 99%, Type 2 90%, Type 3 90%."],
            ["Heart disease prediction", "XGBoost classifier after one-hot encoding categorical clinical fields", "92.93% accuracy on 184 test samples. Heart disease precision 98%, recall 90%."],
        ],
    )
    add_para(doc, "Saved model files include diabetes_rf_model.pkl, diabetes_xgb_final.pkl, diabetes2_xgb_model.pkl, diabetes2_label_encoder.pkl, and heart_xgb_model.pkl.")

    add_heading(doc, "7. Key Risk Factors Identified")
    add_table(
        doc,
        ["Disease Area", "Important Factors"],
        [
            ["Diabetes", "Glucose, BMI, age, insulin, blood pressure, diabetes pedigree/family-history score, pregnancies in the Pima dataset, lifestyle and dietary factors in the large dataset."],
            ["Diabetes type", "Autoantibodies, genetic markers, family history, insulin levels, age, BMI, blood glucose, pancreatic health, steroid use, pregnancy-related fields, lifestyle, smoking, and alcohol."],
            ["Heart disease", "Blocked vessels, exercise-induced chest pain, ST depression/oldpeak, age, sex, cholesterol, blood pressure, fasting blood sugar, maximum heart rate, ECG/stress-test related values."],
        ],
    )

    add_heading(doc, "8. Backend Architecture")
    add_para(doc, "The backend is a Flask API. It loads trained model artifacts on startup, exposes prediction endpoints, handles chat requests, manages hospital data through SQLite repositories, and returns structured JSON to the frontend.")
    add_bullets(
        doc,
        [
            "Prediction modules: predict.py and explain.py.",
            "Simple heart screening: heart_screening.py for non-report users.",
            "Chatbot and RAG: rag.py, intent.py, safety.py, language.py, backend/prompts/system_prompt.txt.",
            "Knowledge indexing: index.py builds ChromaDB from eight knowledge text files.",
            "Database layer: backend/database contains schema, auth, hospital, patient, visit, report, analytics, audit, chat, and summary repositories.",
        ],
    )

    add_heading(doc, "9. RAG Architecture")
    add_para(doc, "The RAG pipeline combines local medical knowledge retrieval with an LLM response layer.")
    add_table(
        doc,
        ["Step", "Description"],
        [
            ["Knowledge creation", "Medical knowledge was written into chunked text files for diabetes, diabetes types, heart disease, general health, foods, Indian diet, South Indian diabetes diet, and Indian heart diet."],
            ["Indexing", "backend/index.py parses [CHUNK_START]/[CHUNK_END] blocks, creates embeddings with all-MiniLM-L6-v2, and stores them in ChromaDB."],
            ["Intent detection", "intent.py detects greeting, thanks, diabetes, heart, diet, prediction, type, and symptom-related queries."],
            ["Retrieval", "The query is expanded and filtered by disease metadata where possible, then relevant chunks are retrieved and compressed."],
            ["Prompting", "The system prompt enforces friendly tone, safety, short formatting, no diagnosis, no exaggerated cure claims, and no internal metadata exposure."],
            ["LLM response", "OpenRouter model nvidia/nemotron-3-nano-omni-30b-a3b-reasoning:free generates the answer when available."],
            ["Fallback", "If the LLM fails, local fallback answers use retrieved chunks or direct rules for common questions."],
            ["Memory and context", "Session memory and patient context help answer follow-up questions without permanently storing private data in the basic browser session flow."],
        ],
    )

    add_heading(doc, "10. Safety and Medical Reliability")
    add_bullets(
        doc,
        [
            "Emergency symptoms such as chest pain and breathing difficulty are handled before LLM generation.",
            "The chatbot always frames output as education and screening, not diagnosis.",
            "Unsafe knowledge wording such as superfood/cure-style claims was cleaned.",
            "Diabetes type is not claimed from the basic diabetes form; it says clinical confirmation is required.",
            "Advanced heart mode is marked as report-based, while simple heart screening is offered for general users.",
            "Sources are retained internally but no longer shown as noisy chunk labels in the UI response.",
        ],
    )

    add_heading(doc, "11. Hospital Login and Role-Based Workflow")
    add_para(doc, "MediRisk was later expanded into a hospital-style workflow with login and role-specific dashboards.")
    add_table(
        doc,
        ["Role", "Main Capabilities"],
        [
            ["Super Admin", "Global hospital analytics, hospital-wise summaries, audit activity, patient insights, portal access monitoring."],
            ["Hospital Admin", "Hospital overview, patient registration, report creation, predictions, patient insights, doctor summaries, portal password reset."],
            ["Doctor", "Patient history, report creation, AI explanations, doctor-written patient summaries, quick predictions."],
            ["Receptionist", "Patient registration, patient search, visit creation, report entry, quick screening."],
            ["Patient", "View own reports, doctor summaries, recommendations, and chatbot with own context."],
        ],
    )
    add_para(doc, "Security note: user and patient passwords are stored as hashes using Werkzeug's generate_password_hash() and verified with check_password_hash(); raw passwords are not stored.")

    add_heading(doc, "12. Database Design")
    add_para(doc, "The SQLite demo database is backend/medirisk.db. The schema supports hospitals, users, patients, patient consents, patient portal accounts, visits, diabetes reports, heart reports, doctor notes, patient-facing summaries, chat sessions/messages, and audit logs. Indexes were added for common patient, visit, account, chat, and analytics lookups.")

    add_heading(doc, "13. Frontend Work")
    add_bullets(
        doc,
        [
            "Landing page: ui/index.html introduces MediRisk, features, stats, and navigation.",
            "Chatbot page: ui/chatbot.html and ui/js/chatbot.js provide chat, quick questions, prediction forms, result cards, session memory id, and theme support.",
            "Hospital pages: hospital-login, super-admin, hospital-admin, doctor, receptionist, and patient pages each have focused role workflows.",
            "Shared role JS and CSS were separated to keep the dashboard maintainable.",
        ],
    )
    add_image_if_exists(doc, "docs/assets/block diagram-full architecture.png", "Full MediRisk architecture diagram.", width=6.3)

    add_heading(doc, "14. Testing and Demo Evidence")
    add_bullets(
        doc,
        [
            "Vector database rebuild completed with 84 chunks.",
            "Prediction scripts loaded models and returned diabetes and heart prediction outputs.",
            "Chat tests covered diabetes symptoms, heart warning signs, heart diet, diabetes diet, weekly diabetes plan, glucose questions, memory recall, and fallback behavior.",
            "Hospital integration tests covered role login, patient registration, portal credential generation, doctor summaries, patient insights, quick prediction, and printing.",
            "Known environment issue: backend must be run through the project virtual environment so xgboost is available.",
        ],
    )

    add_heading(doc, "15. Demo Script for Tomorrow")
    add_table(
        doc,
        ["Demo Step", "What to Show", "Talking Point"],
        [
            ["1", "Open landing page", "Introduce MediRisk as an AI health screening and hospital workflow assistant."],
            ["2", "Open chatbot", "Ask diabetes symptoms and heart diet questions to show RAG guidance."],
            ["3", "Emergency query", "Ask chest pain question to show safety layer before LLM."],
            ["4", "Diabetes prediction", "Enter glucose/BMI/age values and show result card with explanation."],
            ["5", "Simple heart check", "Use symptom-based form for general users."],
            ["6", "Hospital login", "Login as hospital admin/doctor/receptionist/patient and show role-based UI."],
            ["7", "Patient workflow", "Register patient, create visit/report, view insights and doctor summary."],
            ["8", "Architecture", "Explain models, RAG, database, security, and limitations."],
        ],
    )

    add_heading(doc, "16. Limitations and Future Scope")
    add_bullets(
        doc,
        [
            "The system is educational and screening-focused; it is not a diagnostic medical device.",
            "SQLite is suitable for local demo, but PostgreSQL/MySQL would be better for production.",
            "Production use would need JWT/session hardening, encrypted medical records, stricter audit coverage, and deployment secrets management.",
            "Model validation should be expanded with clinically representative datasets before real-world use.",
            "RAG quality can be improved with more guideline-reviewed knowledge, stronger chunk routing, and answer evaluation tests.",
            "Future work can add appointment scheduling, PDF report upload, doctor approval workflow, and full patient chat history UI.",
        ],
    )

    add_heading(doc, "17. Final Outcome")
    add_para(doc, "The final MediRisk project demonstrates an end-to-end AI healthcare system: medical dataset analysis, visual disease-risk exploration, trained ML prediction models, explainable prediction outputs, RAG-based chatbot guidance, multilingual support, emergency safety handling, role-based hospital workflows, patient records, doctor summaries, and patient portal access. It is suitable for an academic demo because it shows both data science depth and application-level integration.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
